import io
import re
from datetime import datetime

from django.conf import settings
from django.core.mail import EmailMessage, get_connection
from docx import Document
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage


def extrair_tags_do_docx(caminho_arquivo):
    """
    Abre o arquivo DOCX e retorna as tags na ordem em que aparecem.
    """
    try:
        doc = Document(caminho_arquivo)
        texto_completo = []

        for paragrafo in doc.paragraphs:
            texto_completo.append(paragrafo.text)

        for tabela in doc.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    texto_completo.append(celula.text)

        conteudo = "\n".join(texto_completo)
        pattern = r'(?:\{\{|\{%\s*if)\s*(\w+)\s*(?:\}\}|%\})'
        tags_encontradas = re.findall(pattern, conteudo)
        return list(dict.fromkeys(tags_encontradas))
    except Exception as exc:  # pragma: no cover - defensive fallback
        print(f"Erro ao ler DOCX: {exc}")
        return []


def carregar_campos_template(template):
    campos = template.configuracao_campos or []
    if campos:
        return campos

    tags = extrair_tags_do_docx(template.arquivo_template.path)
    return [
        {
            'tag': tag,
            'label': tag.replace('_', ' ').title(),
            'tipo': 'text',
            'dependencia': '',
            'opcoes': '',
        }
        for tag in tags
    ]


def renderizar_template_docx(template, dados=None, arquivos=None):
    """
    Renderiza o template DOCX e devolve o arquivo em bytes junto do log.
    """
    dados = dados or {}
    arquivos = arquivos or {}
    campos = carregar_campos_template(template)
    doc = DocxTemplate(template.arquivo_template.path)
    contexto = {}

    for campo in campos:
        tag = campo['tag']
        tipo = campo.get('tipo', 'text')
        valor = dados.get(tag, "")

        if tipo == 'image':
            imagem = arquivos.get(tag)
            contexto[tag] = InlineImage(doc, imagem, width=Mm(160)) if imagem else ""
        elif tipo == 'checkbox':
            if isinstance(valor, str):
                valor = valor.strip().lower() in {'1', 'true', 'on', 'sim', 'yes'}
            contexto[tag] = bool(valor)
        elif tipo == 'date':
            if valor:
                try:
                    valor = datetime.strptime(str(valor), '%Y-%m-%d').strftime('%d/%m/%Y')
                except ValueError:
                    pass
            contexto[tag] = valor or ""
        else:
            contexto[tag] = valor

    doc.render(contexto)
    buffer = io.BytesIO()
    doc.save(buffer)

    conteudo_docx = buffer.getvalue()
    dados_log = {
        chave: str(valor)
        for chave, valor in contexto.items()
        if not isinstance(valor, InlineImage)
    }
    return conteudo_docx, dados_log, campos


def montar_assunto_obrigacao_fazer(pedido):
    if pedido.get('assunto'):
        return pedido['assunto']

    assunto = f"{pedido.get('tipo_demanda_label', 'Solicitacao')} - {pedido.get('cliente', '').strip()}"
    if pedido.get('cpf'):
        assunto += f" - CPF {pedido['cpf']}"
    elif pedido.get('numero_processo'):
        assunto += f" - Processo {pedido['numero_processo']}"
    elif pedido.get('carteira_label'):
        assunto += f" - {pedido['carteira_label']}"
    return assunto.strip(' -')


def montar_corpo_obrigacao_fazer(pedido, solicitante, possui_anexo=False):
    assinatura = solicitante.get_full_name().strip() or solicitante.username

    linhas = [
        "Pessoal, boa tarde.",
        "",
        "Espero que estejam bem.",
        "",
        f"Muito prazer, meu nome e {assinatura} e serei o responsavel pela conducao da carteira de processos relacionada abaixo.",
        "Coloco-me a disposicao para quaisquer esclarecimentos ou necessidades por e-mail, Teams ou WhatsApp.",
        "",
        "Encaminho a solicitacao para cumprimento da demanda abaixo.",
        "",
    ]

    linhas.extend(
        [
            f"Tipo de demanda: {pedido.get('tipo_demanda_label', '')}",
            f"Carteira/Produto: {pedido.get('carteira_label') or pedido.get('carteira') or 'Nao informado'}",
            f"Cliente: {pedido.get('cliente', '')}",
            f"Pedido: {pedido.get('pedido', '')}",
        ]
    )

    if pedido.get('pedido_id'):
        linhas.append(f"Pedido interno: {pedido['pedido_id']}")
    if pedido.get('cpf'):
        linhas.append(f"CPF: {pedido['cpf']}")
    if pedido.get('numero_processo'):
        linhas.append(f"Processo: {pedido['numero_processo']}")

    if pedido.get('tipo_demanda_key') == 'suspensao_liquidacao_reativacao':
        linhas.extend(
            [
                f"Resultado da sentenca: {pedido.get('resultado_sentenca', '')}",
                f"Data da sentenca: {pedido.get('data_sentenca', '')}",
                f"ID do oficio: {pedido.get('id_oficio', '')}",
                f"Comarca: {pedido.get('comarca', '')}",
            ]
        )

    if pedido.get('tipo_demanda_key') == 'cobranca_consignado' and pedido.get('observacoes'):
        linhas.extend(
            [
                "",
                "Detalhamento para emissao de boletos:",
                pedido['observacoes'],
            ]
        )

    if pedido.get('tipo_demanda_key') == 'cartoes' and pedido.get('observacoes'):
        linhas.extend(
            [
                "",
                "Detalhamento da solicitacao de cartoes:",
                pedido['observacoes'],
            ]
        )

    if pedido.get('observacoes') and pedido.get('tipo_demanda_key') not in {'cobranca_consignado', 'cartoes'}:
        linhas.extend(
            [
                "",
                "Observacoes:",
                pedido['observacoes'],
            ]
        )

    if possui_anexo:
        linhas.extend(
            [
                "",
                "Segue em anexo o PDF correspondente a este pedido.",
            ]
        )

    linhas.extend(
        [
            "",
            "Fico a disposicao para eventuais esclarecimentos.",
            "",
            "Atenciosamente,",
            assinatura,
        ]
    )
    return "\n".join(linhas)


def enviar_lote_obrigacao_fazer(pedidos, solicitante):
    enviados = []
    falhas = []
    connection = get_connection(fail_silently=False)
    connection.open()

    try:
        for pedido in pedidos:
            identificador = (
                pedido.get('identificador')
                or pedido.get('pedido_id')
                or f"linha {pedido.get('linha_origem', '?')}"
            )

            try:
                anexo_nome = pedido.get('anexo_pdf_nome')
                anexo_conteudo = pedido.get('anexo_pdf_conteudo')
                email = EmailMessage(
                    subject=montar_assunto_obrigacao_fazer(pedido),
                    body=montar_corpo_obrigacao_fazer(
                        pedido,
                        solicitante=solicitante,
                        possui_anexo=bool(anexo_conteudo),
                    ),
                    from_email=settings.DEFAULT_FROM_EMAIL,
                    to=[pedido['email_destinatario']],
                    cc=pedido.get('cc_list') or None,
                    connection=connection,
                )
                if anexo_conteudo:
                    email.attach(anexo_nome, anexo_conteudo, 'application/pdf')

                email.send(fail_silently=False)
                enviados.append(identificador)
            except Exception as exc:  # pragma: no cover - network/backend variability
                falhas.append(
                    {
                        'identificador': identificador,
                        'linha': pedido.get('linha_origem'),
                        'erro': str(exc),
                    }
                )
    finally:
        connection.close()

    return enviados, falhas
