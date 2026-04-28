import csv
import hashlib
import io
import re
import unicodedata
from pathlib import Path

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import connection, transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    ArquivoProcesso,
    ExecucaoCapturaProcesso,
    LoteImportacaoProcessos,
    Processo,
    ProcessoStatus,
    TipoArquivoProcesso,
)
from .processos_constants import TRIBUNAIS_ESTADUAIS


STATUSS_FINAIS_PROCESSO = {
    ProcessoStatus.BAIXADO,
    ProcessoStatus.PROCESSO_NAO_ENCONTRADO,
    ProcessoStatus.PETICAO_NAO_LOCALIZADA,
    ProcessoStatus.FALHA_TECNICA,
}

STATUSS_ABERTOS_EXECUCAO = {
    ProcessoStatus.PENDENTE,
    ProcessoStatus.EM_FILA,
    ProcessoStatus.PROCESSANDO,
}


_COLUMN_ALIASES_PROCESSOS = {
    'numero_processo': [
        'numero_processo',
        'numero_cnj',
        'processo',
        'cnj',
    ],
    'usuario_email': [
        'usuario_email',
        'email_usuario',
        'usuario',
        'email',
    ],
    'observacao': [
        'observacao',
        'observacoes',
        'obs',
    ],
    'referencia_interna': [
        'referencia_interna',
        'referencia',
        'codigo_interno',
        'protocolo',
    ],
}

def _normalizar_texto(valor):
    valor = unicodedata.normalize('NFKD', str(valor or ''))
    valor = valor.encode('ascii', 'ignore').decode('ascii')
    valor = valor.strip().lower()
    valor = re.sub(r'[^a-z0-9]+', '_', valor)
    return valor.strip('_')


_ALIAS_MAP_PROCESSOS = {}
for _campo, _aliases in _COLUMN_ALIASES_PROCESSOS.items():
    for _alias in _aliases:
        _ALIAS_MAP_PROCESSOS[_normalizar_texto(_alias)] = _campo


def _normalizar_celula(valor):
    if valor is None:
        return ''
    return str(valor).strip()


def _carregar_planilha_csv(conteudo):
    ultimo_erro = None
    texto = None
    for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            texto = conteudo.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            ultimo_erro = exc

    if texto is None:
        raise ValidationError(f"Nao foi possivel ler o CSV enviado: {ultimo_erro}")

    leitor = csv.DictReader(io.StringIO(texto))
    if not leitor.fieldnames:
        raise ValidationError("A planilha CSV precisa ter um cabecalho na primeira linha.")

    return list(leitor.fieldnames), list(leitor)


def _carregar_planilha_xlsx(conteudo):
    workbook = load_workbook(filename=io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        linhas = worksheet.iter_rows(values_only=True)
        try:
            cabecalhos = next(linhas)
        except StopIteration as exc:
            raise ValidationError("A planilha XLSX enviada esta vazia.") from exc

        cabecalhos = [_normalizar_celula(valor) for valor in cabecalhos]
        registros = [dict(zip(cabecalhos, linha)) for linha in linhas]
        return cabecalhos, registros
    finally:
        workbook.close()


def _identificar_colunas(cabecalhos):
    mapa_colunas = {}
    for cabecalho in cabecalhos:
        campo = _ALIAS_MAP_PROCESSOS.get(_normalizar_texto(cabecalho))
        if campo and campo not in mapa_colunas:
            mapa_colunas[campo] = cabecalho
    return mapa_colunas


def _formatar_numero_cnj(digitos):
    return f"{digitos[:7]}-{digitos[7:9]}.{digitos[9:13]}.{digitos[13]}.{digitos[14:16]}.{digitos[16:20]}"


def _validar_digitos_cnj(digitos):
    numero_sem_dv = digitos[:7] + digitos[9:]
    dv_informado = int(digitos[7:9])
    dv_calculado = 98 - (int(numero_sem_dv + '00') % 97)
    return dv_calculado == dv_informado


def normalizar_numero_cnj(numero_processo):
    digitos = re.sub(r'[^0-9]', '', str(numero_processo or ''))
    if len(digitos) != 20:
        raise ValidationError("Informe um numero CNJ valido com 20 digitos.")
    if not _validar_digitos_cnj(digitos):
        raise ValidationError("O numero do processo nao passou na validacao do CNJ.")
    return _formatar_numero_cnj(digitos)


def resolver_tribunal_por_cnj(numero_cnj):
    digitos = re.sub(r'[^0-9]', '', numero_cnj or '')
    justica = digitos[13:14]
    tribunal_numero = digitos[14:16]
    if justica != '8':
        raise ValidationError("O processo informado nao pertence a um TJ estadual da lista suportada.")

    codigo = f"8.{tribunal_numero}"
    nome = TRIBUNAIS_ESTADUAIS.get(codigo)
    if not nome:
        raise ValidationError(f"Nao existe mapeamento de tribunal para o codigo {codigo}.")
    return codigo, nome


def _localizar_usuario_por_email(email):
    email = (email or '').strip().lower()
    if not email:
        raise ValidationError("Informe o usuario_email.")

    usuario = User.objects.filter(
        is_active=True,
    ).filter(
        Q(username__iexact=email) | Q(email__iexact=email),
    ).first()
    if not usuario:
        raise ValidationError(f"Nao existe usuario ativo no DOC para o email '{email}'.")
    return usuario


def parsear_planilha_processos(upload):
    extensao = Path(upload.name or '').suffix.lower()
    conteudo = upload.read()
    upload.seek(0)

    if extensao == '.csv':
        cabecalhos, registros = _carregar_planilha_csv(conteudo)
    elif extensao == '.xlsx':
        cabecalhos, registros = _carregar_planilha_xlsx(conteudo)
    else:
        raise ValidationError("Envie uma planilha no formato .xlsx ou .csv.")

    mapa_colunas = _identificar_colunas(cabecalhos)
    faltantes = [
        campo for campo in ('numero_processo', 'usuario_email')
        if campo not in mapa_colunas
    ]
    if faltantes:
        raise ValidationError("Colunas obrigatorias ausentes: " + ", ".join(faltantes) + ".")

    linhas = []
    erros = []
    vistos = set()
    total_linhas = 0

    for indice, registro in enumerate(registros, start=2):
        linha = {}
        for campo in _COLUMN_ALIASES_PROCESSOS:
            cabecalho = mapa_colunas.get(campo)
            linha[campo] = _normalizar_celula(registro.get(cabecalho, '')) if cabecalho else ''

        if not any(linha.values()):
            continue

        total_linhas += 1

        if not linha['numero_processo'] or not linha['usuario_email']:
            erros.append(
                f"Linha {indice}: faltam valores para numero_processo e/ou usuario_email."
            )
            continue

        try:
            usuario = _localizar_usuario_por_email(linha['usuario_email'])
            numero_cnj = normalizar_numero_cnj(linha['numero_processo'])
            tribunal_codigo, tribunal_nome = resolver_tribunal_por_cnj(numero_cnj)
        except ValidationError as exc:
            erros.extend([f"Linha {indice}: {mensagem}" for mensagem in exc.messages])
            continue

        chave_duplicidade = (usuario.id, numero_cnj)
        if chave_duplicidade in vistos:
            erros.append(
                f"Linha {indice}: o processo {numero_cnj} esta duplicado para o usuario {usuario.username}."
            )
            continue
        vistos.add(chave_duplicidade)

        linhas.append(
            {
                'linha_origem': indice,
                'usuario': usuario,
                'usuario_email': linha['usuario_email'].strip().lower(),
                'numero_cnj': numero_cnj,
                'tribunal_codigo': tribunal_codigo,
                'tribunal_nome': tribunal_nome,
                'observacao': linha['observacao'],
                'referencia_interna': linha['referencia_interna'],
                'existe_no_banco': Processo.objects.filter(
                    usuario=usuario,
                    numero_cnj=numero_cnj,
                ).exists(),
            }
        )

    if total_linhas == 0 and not erros:
        raise ValidationError("Nenhuma linha preenchida foi encontrada na planilha.")

    if erros:
        raise ValidationError(erros)

    return {
        'linhas': linhas,
        'total_linhas': total_linhas,
        'total_validas': len(linhas),
        'total_invalidas': total_linhas - len(linhas),
    }


def enqueue_processo_para_captura(processo, lote=None, mensagem=''):
    execucao_aberta = processo.execucoes.filter(
        status__in=STATUSS_ABERTOS_EXECUCAO,
    ).order_by('-criado_em').first()

    if execucao_aberta:
        atualizacoes = []
        if lote and execucao_aberta.lote_id is None:
            execucao_aberta.lote = lote
            atualizacoes.append('lote')
        if mensagem:
            execucao_aberta.mensagem = mensagem
            atualizacoes.append('mensagem')
        if atualizacoes:
            execucao_aberta.save(update_fields=atualizacoes)

        processo.status_atual = ProcessoStatus.EM_FILA
        processo.ultima_mensagem = mensagem or processo.ultima_mensagem
        processo.save(update_fields=['status_atual', 'ultima_mensagem', 'atualizado_em'])
        return execucao_aberta

    execucao = ExecucaoCapturaProcesso.objects.create(
        processo=processo,
        lote=lote,
        status=ProcessoStatus.EM_FILA,
        mensagem=mensagem,
    )
    processo.status_atual = ProcessoStatus.EM_FILA
    processo.ultima_mensagem = mensagem
    processo.save(update_fields=['status_atual', 'ultima_mensagem', 'atualizado_em'])
    return execucao


@transaction.atomic
def importar_lote_processos(usuario_admin, planilha, dados_planilha):
    planilha.seek(0)
    lote = LoteImportacaoProcessos.objects.create(
        usuario=usuario_admin,
        arquivo_planilha=planilha,
        total_linhas=dados_planilha['total_linhas'],
        total_validas=dados_planilha['total_validas'],
        total_invalidas=dados_planilha['total_invalidas'],
    )

    criados = 0
    atualizados = 0
    processos_ids = []

    for linha in dados_planilha['linhas']:
        processo, criado = Processo.objects.update_or_create(
            usuario=linha['usuario'],
            numero_cnj=linha['numero_cnj'],
            defaults={
                'tribunal_codigo': linha['tribunal_codigo'],
                'tribunal_nome': linha['tribunal_nome'],
                'observacao': linha['observacao'],
                'referencia_interna': linha['referencia_interna'],
            },
        )
        if criado:
            criados += 1
        else:
            atualizados += 1

        mensagem = (
            f"Processo importado no lote #{lote.pk} "
            f"(linha {linha['linha_origem']})."
        )
        enqueue_processo_para_captura(processo, lote=lote, mensagem=mensagem)
        processos_ids.append(processo.id)

    lote.total_processos_criados = criados
    lote.total_processos_atualizados = atualizados
    lote.resultado_importacao = {
        'processos_ids': processos_ids,
        'linhas_validas': [
            {
                'linha_origem': linha['linha_origem'],
                'usuario_email': linha['usuario_email'],
                'numero_cnj': linha['numero_cnj'],
            }
            for linha in dados_planilha['linhas']
        ],
    }
    lote.save(
        update_fields=[
            'total_processos_criados',
            'total_processos_atualizados',
            'resultado_importacao',
        ]
    )
    return lote


@transaction.atomic
def reprocessar_processo(processo, usuario=None):
    mensagem = "Reprocessamento solicitado manualmente."
    if usuario:
        mensagem = f"{mensagem} Responsavel: {usuario.username}."
    return enqueue_processo_para_captura(processo, mensagem=mensagem)


@transaction.atomic
def claim_execucao_para_worker(worker_id):
    queryset = ExecucaoCapturaProcesso.objects.select_related('processo', 'processo__usuario').filter(
        status__in=[ProcessoStatus.PENDENTE, ProcessoStatus.EM_FILA]
    ).order_by('criado_em', 'id')

    if connection.features.has_select_for_update_skip_locked:
        queryset = queryset.select_for_update(skip_locked=True)
    else:
        queryset = queryset.select_for_update()

    execucao = queryset.first()
    if not execucao:
        return None

    execucao.status = ProcessoStatus.PROCESSANDO
    execucao.worker_id = worker_id
    execucao.tentativa += 1
    execucao.iniciada_em = timezone.now()
    execucao.heartbeat_em = execucao.iniciada_em
    execucao.finalizada_em = None
    execucao.save(
        update_fields=[
            'status',
            'worker_id',
            'tentativa',
            'iniciada_em',
            'heartbeat_em',
            'finalizada_em',
            'atualizado_em',
        ]
    )

    execucao.processo.status_atual = ProcessoStatus.PROCESSANDO
    execucao.processo.ultima_mensagem = "Captura em andamento pelo worker."
    execucao.processo.save(update_fields=['status_atual', 'ultima_mensagem', 'atualizado_em'])
    return execucao


def registrar_heartbeat_execucao(execucao, mensagem=''):
    execucao.heartbeat_em = timezone.now()
    if mensagem:
        execucao.mensagem = mensagem
    execucao.save(update_fields=['heartbeat_em', 'mensagem', 'atualizado_em'])
    return execucao


@transaction.atomic
def concluir_execucao_com_arquivo(execucao, arquivo_upload, checksum='', mensagem='', tribunal_url=''):
    conteudo = arquivo_upload.read()
    arquivo_upload.seek(0)

    checksum = checksum or hashlib.sha256(conteudo).hexdigest()

    ArquivoProcesso.objects.filter(
        processo=execucao.processo,
        tipo=TipoArquivoProcesso.PETICAO_INICIAL,
        atual=True,
    ).update(atual=False)

    arquivo = ArquivoProcesso.objects.create(
        processo=execucao.processo,
        execucao=execucao,
        tipo=TipoArquivoProcesso.PETICAO_INICIAL,
        checksum=checksum,
        atual=True,
    )
    arquivo.arquivo.save(
        Path(arquivo_upload.name or 'peticao_inicial.pdf').name,
        ContentFile(conteudo),
        save=True,
    )

    execucao.status = ProcessoStatus.BAIXADO
    execucao.finalizada_em = timezone.now()
    execucao.heartbeat_em = execucao.finalizada_em
    execucao.mensagem = mensagem or "Peticao inicial capturada com sucesso."
    if tribunal_url:
        execucao.tribunal_url = tribunal_url
    execucao.save(
        update_fields=[
            'status',
            'finalizada_em',
            'heartbeat_em',
            'mensagem',
            'tribunal_url',
            'atualizado_em',
        ]
    )

    execucao.processo.status_atual = ProcessoStatus.BAIXADO
    execucao.processo.ultima_mensagem = execucao.mensagem
    execucao.processo.ultimo_sucesso_em = execucao.finalizada_em
    execucao.processo.save(
        update_fields=['status_atual', 'ultima_mensagem', 'ultimo_sucesso_em', 'atualizado_em']
    )
    return arquivo


@transaction.atomic
def falhar_execucao(execucao, status_final, mensagem, tribunal_url=''):
    if status_final not in STATUSS_FINAIS_PROCESSO - {ProcessoStatus.BAIXADO}:
        raise ValidationError("Status final invalido para falha de execucao.")

    execucao.status = status_final
    execucao.finalizada_em = timezone.now()
    execucao.heartbeat_em = execucao.finalizada_em
    execucao.mensagem = mensagem
    if tribunal_url:
        execucao.tribunal_url = tribunal_url
    execucao.save(
        update_fields=[
            'status',
            'finalizada_em',
            'heartbeat_em',
            'mensagem',
            'tribunal_url',
            'atualizado_em',
        ]
    )

    execucao.processo.status_atual = status_final
    execucao.processo.ultima_mensagem = mensagem
    execucao.processo.save(update_fields=['status_atual', 'ultima_mensagem', 'atualizado_em'])
    return execucao
