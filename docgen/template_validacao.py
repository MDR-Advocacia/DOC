"""Pré-validação de templates DOCX no upload.

Objetivo: nenhum modelo com tag quebrada entra no catálogo. Antes de salvar,
o arquivo é varrido e cada problema é devolvido *localizado* (onde está, qual
é a tag, o que fazer) para o usuário corrigir no Word.

Por que isso existe: erros como ``{{ e-mail }}`` são **sintaxe Jinja válida**
— o parser lê como a subtração ``e - mail``. Passavam na validação antiga e só
explodiam na hora de gerar o documento, com a mensagem inútil
``'e' is undefined``, sem dizer qual tag nem onde.

Sobre "página": o Word calcula a paginação ao renderizar; o .docx não guarda
isso. Só dá para estimar quando o arquivo tem marcadores de quebra
(``lastRenderedPageBreak`` ou quebra manual). Quando não tem, ``pagina`` vem
``None`` e o usuário localiza pelo número do parágrafo + o trecho do texto.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field

from docx import Document


# Conteúdo entre {{ }} ou {% %}, e a tag inteira.
TAG_PATTERN = re.compile(r'\{\{[\s\S]*?\}\}|\{%[\s\S]*?%\}')
IDENTIFICADOR_LIMPO = re.compile(r'^[A-Za-z_][A-Za-z0-9_]*$')

# Keywords/operadores que indicam expressão Jinja legítima — quando aparecem,
# não tratamos "duas palavras separadas por espaço" como erro de digitação.
KEYWORDS_JINJA = {
    'if', 'elif', 'else', 'endif', 'for', 'endfor', 'in', 'is', 'not',
    'and', 'or', 'true', 'false', 'none', 'True', 'False', 'None',
}
BLOCOS_ABRE = {'if': 'endif', 'for': 'endfor'}


@dataclass
class ProblemaTemplate:
    """Um problema encontrado, já localizado no documento."""

    tipo: str            # slug estável: 'hifen', 'espaco', 'ponto', ...
    titulo: str          # rótulo curto pro usuário
    explicacao: str      # por que quebra
    tag: str             # a tag como está no documento
    sugestao: str        # como deveria ficar ('' quando não dá pra inferir)
    local: str           # 'corpo' | 'tabela' | 'cabeçalho' | 'rodapé'
    paragrafo: int | None = None
    pagina: int | None = None
    trecho: str = ''
    gravidade: str = 'erro'   # 'erro' bloqueia o upload; 'aviso' não


@dataclass
class RelatorioValidacao:
    problemas: list[ProblemaTemplate] = field(default_factory=list)
    erro_tecnico: str = ''
    paginacao_disponivel: bool = False

    @property
    def erros(self) -> list[ProblemaTemplate]:
        return [p for p in self.problemas if p.gravidade == 'erro']

    @property
    def avisos(self) -> list[ProblemaTemplate]:
        return [p for p in self.problemas if p.gravidade == 'aviso']

    @property
    def ok(self) -> bool:
        return not self.erros and not self.erro_tecnico


# ---------------------------------------------------------------------------
# Catálogo de erros de tag
# ---------------------------------------------------------------------------

def _tem_acento(texto: str) -> bool:
    return any(unicodedata.category(c) == 'Mn' for c in unicodedata.normalize('NFD', texto))


def _sem_acento(texto: str) -> str:
    nfd = unicodedata.normalize('NFD', texto)
    return ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')


def _classificar_conteudo(miolo: str, tag_completa: str) -> list[dict]:
    """Recebe o miolo de uma tag ``{{ ... }}`` e devolve os problemas achados."""
    achados: list[dict] = []
    conteudo = miolo.strip()

    if not conteudo:
        return [{
            'tipo': 'vazia',
            'titulo': 'Tag vazia',
            'explicacao': 'Existe um `{{ }}` sem nada dentro. O sistema não tem o que preencher aqui.',
            'sugestao': '',
        }]

    # Expressão Jinja de verdade (filtro, comparação, atributo de objeto):
    # não é o caso de uso deste sistema, mas também não é erro de digitação.
    tem_operador = bool(re.search(r'[=<>!|~+*/]|\bis\b|\bin\b|\band\b|\bor\b|\bnot\b', conteudo))

    if '-' in conteudo:
        sugerido = conteudo.replace('-', '_')
        achados.append({
            'tipo': 'hifen',
            'titulo': 'Hífen dentro da tag',
            'explicacao': (
                'O hífen é lido como sinal de subtração. Em `{{ e-mail }}` o sistema '
                'procura uma variável chamada `e` e outra chamada `mail`, não acha, '
                'e a geração falha com "\'e\' is undefined". '
                'Nome de campo só aceita letras, números e `_`.'
            ),
            'sugestao': '{{ %s }}' % sugerido,
        })

    if '.' in conteudo and not tem_operador:
        sugerido = conteudo.replace('.', '')
        achados.append({
            'tipo': 'ponto',
            'titulo': 'Ponto dentro da tag',
            'explicacao': (
                'O ponto é lido como acesso a atributo. Em `{{ N.º_CNJ }}` o sistema '
                'procura a variável `N` e nela um campo `º_CNJ`, e a geração falha.'
            ),
            'sugestao': '{{ %s }}' % sugerido,
        })

    partes = conteudo.split()
    if len(partes) > 1 and not tem_operador and not (set(partes) & KEYWORDS_JINJA):
        sugerido = '_'.join(partes)
        achados.append({
            'tipo': 'espaco',
            'titulo': 'Espaço no meio do nome do campo',
            'explicacao': (
                'Nome de campo não pode ter espaço — o Jinja para de ler no primeiro '
                'espaço e acusa erro de sintaxe. Use `_` no lugar.'
            ),
            'sugestao': '{{ %s }}' % sugerido,
        })

    if re.match(r'^\d', conteudo):
        achados.append({
            'tipo': 'comeca_com_numero',
            'titulo': 'Nome do campo começa com número',
            'explicacao': (
                'Nome de campo não pode começar com dígito. `{{ 2_via }}` é inválido; '
                'use `{{ via_2 }}` ou `{{ segunda_via }}`.'
            ),
            'sugestao': '{{ campo_%s }}' % conteudo.replace(' ', '_'),
        })

    if _tem_acento(conteudo) or 'ç' in conteudo.lower():
        achados.append({
            'tipo': 'acento',
            'titulo': 'Acento ou cedilha no nome do campo',
            'explicacao': (
                'O Word costuma fragmentar palavras acentuadas no XML e a tag chega '
                'quebrada no gerador. O sistema renomeia sozinho no upload, mas o '
                'ideal é já subir sem acento.'
            ),
            'sugestao': '{{ %s }}' % _sem_acento(conteudo).replace('ç', 'c').replace('Ç', 'C'),
            'gravidade': 'aviso',
        })

    # Qualquer outro caractere fora de [A-Za-z0-9_ ] que ainda não foi coberto.
    if not tem_operador:
        estranhos = set(re.findall(r'[^\w\s]', conteudo, flags=re.UNICODE)) - {'-', '.'}
        if estranhos:
            achados.append({
                'tipo': 'caractere_invalido',
                'titulo': 'Caractere inválido no nome do campo',
                'explicacao': (
                    'Encontrei %s dentro da tag. Nome de campo só aceita letras, '
                    'números e `_`.' % ', '.join(f'`{c}`' for c in sorted(estranhos))
                ),
                'sugestao': '{{ %s }}' % re.sub(r'[^\w]', '_', conteudo),
            })

    return achados


def _classificar_tag(tag: str) -> list[dict]:
    """Classifica uma tag inteira (com as chaves)."""
    if tag.startswith('{{'):
        return _classificar_conteudo(tag[2:-2], tag)

    # Bloco {% ... %}: valida só a keyword e o nome da variável do if/for.
    miolo = tag[2:-2].strip()
    partes = miolo.split()
    if not partes:
        return [{
            'tipo': 'vazia',
            'titulo': 'Bloco vazio',
            'explicacao': 'Existe um `{% %}` sem instrução dentro.',
            'sugestao': '',
        }]

    keyword = partes[0]
    conhecidas = {
        'if', 'elif', 'else', 'endif', 'for', 'endfor', 'set', 'endset',
        'macro', 'endmacro', 'block', 'endblock', 'raw', 'endraw',
    }
    if keyword not in conhecidas:
        return [{
            'tipo': 'bloco_desconhecido',
            'titulo': f'Instrução desconhecida: `{keyword}`',
            'explicacao': (
                f'`{{% {keyword} %}}` não é um comando válido. Verifique se não foi '
                'digitado errado — `endif` virando `endiv` é o caso mais comum.'
            ),
            'sugestao': '',
        }]

    if keyword in {'if', 'elif'} and len(partes) > 1:
        return _classificar_conteudo(' '.join(partes[1:]), tag)

    return []


# ---------------------------------------------------------------------------
# Varredura do documento
# ---------------------------------------------------------------------------

def _mapa_paginas(doc) -> tuple[list[int | None], bool]:
    """Estima a página de cada parágrafo do corpo.

    Só funciona quando o .docx tem marcadores de quebra. O Word grava
    ``lastRenderedPageBreak`` com a paginação do último render — quando existe,
    é o sinal mais fiel. Sem nenhum marcador, devolve None em tudo, porque
    chutar página por contagem de caracteres erraria em documento longo.
    """
    xmls = [p._p.xml for p in doc.paragraphs]
    usa_render = any('lastRenderedPageBreak' in x for x in xmls)
    padrao_manual = re.compile(r'<w:br[^>]*w:type="page"')

    if not usa_render and not any(padrao_manual.search(x) for x in xmls):
        return [None] * len(xmls), False

    mapa: list[int | None] = []
    pagina = 1
    for xml in xmls:
        mapa.append(pagina)
        if usa_render:
            pagina += xml.count('lastRenderedPageBreak')
        else:
            pagina += len(padrao_manual.findall(xml))
    return mapa, True


def _trecho(texto: str, inicio: int, fim: int, margem: int = 45) -> str:
    ini = max(0, inicio - margem)
    prefixo = '…' if ini > 0 else ''
    sufixo = '…' if fim + margem < len(texto) else ''
    return f"{prefixo}{texto[ini:min(len(texto), fim + margem)]}{sufixo}"


def _blocos_de_texto(doc):
    """Gera (texto, local, indice_paragrafo) de todo o documento."""
    for i, p in enumerate(doc.paragraphs, start=1):
        if p.text:
            yield p.text, 'corpo', i

    for t_idx, tabela in enumerate(doc.tables, start=1):
        for l_idx, linha in enumerate(tabela.rows, start=1):
            for c_idx, celula in enumerate(linha.cells, start=1):
                for p in celula.paragraphs:
                    if p.text:
                        yield p.text, f'tabela {t_idx} (linha {l_idx}, coluna {c_idx})', None

    for s_idx, section in enumerate(doc.sections, start=1):
        for p in section.header.paragraphs:
            if p.text:
                yield p.text, f'cabeçalho (seção {s_idx})', None
        for p in section.footer.paragraphs:
            if p.text:
                yield p.text, f'rodapé (seção {s_idx})', None


def _checar_chaves_desbalanceadas(texto: str, local: str, paragrafo, pagina):
    """Detecta `{{` sem `}}`, `{%` sem `%}` e chave simples parecendo tag."""
    problemas = []

    sem_tags = TAG_PATTERN.sub('', texto)

    if '{{' in sem_tags or '}}' in sem_tags:
        problemas.append(ProblemaTemplate(
            tipo='chave_desbalanceada',
            titulo='Chave aberta e não fechada',
            explicacao=(
                'Sobrou um `{{` sem o `}}` correspondente (ou o contrário) neste '
                'parágrafo. Tudo daí em diante deixa de ser lido como texto.'
            ),
            tag=sem_tags.strip()[:80],
            sugestao='',
            local=local, paragrafo=paragrafo, pagina=pagina,
            trecho=_trecho(texto, 0, min(len(texto), 90)),
        ))

    if ('{%' in sem_tags) != ('%}' in sem_tags) and ('{%' in sem_tags or '%}' in sem_tags):
        problemas.append(ProblemaTemplate(
            tipo='bloco_desbalanceado',
            titulo='Bloco `{% %}` aberto e não fechado',
            explicacao='Sobrou um `{%` sem o `%}` correspondente neste parágrafo.',
            tag=sem_tags.strip()[:80],
            sugestao='',
            local=local, paragrafo=paragrafo, pagina=pagina,
            trecho=_trecho(texto, 0, min(len(texto), 90)),
        ))

    # `{ nome }` — usuário esqueceu a segunda chave. Chave simples é rara em
    # peça jurídica, então o falso positivo é aceitável como aviso.
    for m in re.finditer(r'(?<![{}])\{\s*([A-Za-z_][\w\s]{1,40})\s*\}(?![{}])', sem_tags):
        problemas.append(ProblemaTemplate(
            tipo='chave_simples',
            titulo='Chave simples — faltou dobrar',
            explicacao=(
                'Isso parece um campo escrito com uma chave só. O sistema só '
                'reconhece campos com chave dupla: `{{ campo }}`.'
            ),
            tag=m.group(0),
            sugestao='{{ %s }}' % m.group(1).strip().replace(' ', '_'),
            local=local, paragrafo=paragrafo, pagina=pagina,
            trecho=_trecho(sem_tags, m.start(), m.end()),
            gravidade='aviso',
        ))

    return problemas


def _checar_blocos_abertos(todas_tags: list[str]) -> list[ProblemaTemplate]:
    """Confere se todo `{% if %}` / `{% for %}` tem seu fechamento."""
    pilha: list[str] = []
    problemas: list[ProblemaTemplate] = []

    for tag in todas_tags:
        if not tag.startswith('{%'):
            continue
        partes = tag[2:-2].strip().split()
        if not partes:
            continue
        kw = partes[0]
        if kw in BLOCOS_ABRE:
            pilha.append(kw)
        elif kw in {'endif', 'endfor'}:
            esperado = 'if' if kw == 'endif' else 'for'
            if pilha and pilha[-1] == esperado:
                pilha.pop()
            else:
                problemas.append(ProblemaTemplate(
                    tipo='fechamento_sobrando',
                    titulo=f'`{{% {kw} %}}` sem abertura correspondente',
                    explicacao=(
                        f'Achei um `{{% {kw} %}}` que não fecha nenhum '
                        f'`{{% {esperado} %}}` aberto. Provavelmente sobrou de uma '
                        'edição anterior.'
                    ),
                    tag=tag, sugestao='', local='documento',
                ))

    for kw in pilha:
        problemas.append(ProblemaTemplate(
            tipo='bloco_nao_fechado',
            titulo=f'`{{% {kw} %}}` aberto e nunca fechado',
            explicacao=(
                f'Todo `{{% {kw} %}}` precisa de um `{{% {BLOCOS_ABRE[kw]} %}}` '
                'em algum ponto do documento.'
            ),
            tag='{%% %s ... %%}' % kw, sugestao='{%% %s %%}' % BLOCOS_ABRE[kw],
            local='documento',
        ))

    return problemas


def analisar_template_docx(docx_bytes: bytes) -> RelatorioValidacao:
    """Varre o .docx e devolve todos os problemas de tag, já localizados."""
    relatorio = RelatorioValidacao()

    try:
        doc = Document(io.BytesIO(docx_bytes))
    except Exception as exc:
        relatorio.erro_tecnico = f"Não foi possível abrir o arquivo: {exc}"
        return relatorio

    paginas, tem_paginacao = _mapa_paginas(doc)
    relatorio.paginacao_disponivel = tem_paginacao

    todas_tags: list[str] = []
    vistos: set[tuple] = set()

    for texto, local, idx_paragrafo in _blocos_de_texto(doc):
        pagina = (
            paginas[idx_paragrafo - 1]
            if idx_paragrafo and idx_paragrafo <= len(paginas)
            else None
        )

        for m in TAG_PATTERN.finditer(texto):
            tag = m.group()
            todas_tags.append(tag)
            for achado in _classificar_tag(tag):
                chave = (achado['tipo'], tag, local, idx_paragrafo)
                if chave in vistos:
                    continue
                vistos.add(chave)
                relatorio.problemas.append(ProblemaTemplate(
                    tipo=achado['tipo'],
                    titulo=achado['titulo'],
                    explicacao=achado['explicacao'],
                    tag=tag,
                    sugestao=achado.get('sugestao', ''),
                    local=local,
                    paragrafo=idx_paragrafo,
                    pagina=pagina,
                    trecho=_trecho(texto, m.start(), m.end()),
                    gravidade=achado.get('gravidade', 'erro'),
                ))

        relatorio.problemas.extend(
            _checar_chaves_desbalanceadas(texto, local, idx_paragrafo, pagina)
        )

    relatorio.problemas.extend(_checar_blocos_abertos(todas_tags))

    # Rede de segurança: mesmo sem cair em nenhuma regra acima, o documento
    # tem que parsear e renderizar. Pega qualquer coisa que o catálogo não
    # previu — e é o que garante que nada quebrado passe.
    if not relatorio.erros:
        erro = _simular_geracao(docx_bytes)
        if erro:
            relatorio.erro_tecnico = erro

    return relatorio


def _simular_geracao(docx_bytes: bytes) -> str:
    """Renderiza o template com todos os campos vazios, igual a geração real.

    Devolve '' se passou, ou a mensagem traduzida do erro.
    """
    from docxtpl import DocxTemplate
    from jinja2.exceptions import TemplateSyntaxError, UndefinedError

    try:
        tpl = DocxTemplate(io.BytesIO(docx_bytes))
        variaveis = tpl.get_undeclared_template_variables()
    except TemplateSyntaxError as exc:
        return _mensagem_sintaxe(str(exc.message))
    except Exception as exc:
        return f"Não foi possível ler o modelo: {exc}"

    try:
        tpl.render({v: "" for v in variaveis})
    except UndefinedError as exc:
        return (
            f"A geração falharia com: \"{exc}\". Isso quase sempre é uma tag com "
            "hífen ou ponto no nome, que o sistema lê como conta matemática."
        )
    except Exception as exc:
        return f"A geração falharia com: {exc}"

    return ""


def _mensagem_sintaxe(msg_cru: str) -> str:
    traducoes = [
        (
            re.compile(r"expected token 'end of (?:print|statement) block', got '(\S+)'", re.I),
            lambda m: (
                f"Uma tag tem espaço onde deveria ter `_`, perto de `{m.group(1)}`. "
                "Ex.: `{{ endereco eletronico }}` → `{{ endereco_eletronico }}`."
            ),
        ),
        (
            re.compile(r"unexpected end of template", re.I),
            lambda m: "Um `{% if %}` ou `{% for %}` ficou aberto sem o fechamento.",
        ),
        (
            re.compile(r"Encountered unknown tag '(\w+)'", re.I),
            lambda m: f"Instrução desconhecida `{{% {m.group(1)} %}}` — verifique a digitação.",
        ),
    ]
    for pattern, builder in traducoes:
        m = pattern.search(msg_cru)
        if m:
            return builder(m) + f" (erro técnico: {msg_cru})"
    return f"O Word quebrou uma tag do modelo. Erro técnico: {msg_cru}"
