import csv
import io
import re
import unicodedata
from pathlib import Path

from django import forms
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook


COLUNAS_OBRIGATORIAS = [
    {
        'nome': 'cliente',
        'descricao': 'Nome do cliente relacionado ao pedido.',
    },
    {
        'nome': 'tipo_demanda',
        'descricao': 'Use valores como suspensao_liquidacao_reativacao, cobranca_consignado ou cartoes.',
    },
    {
        'nome': 'pedido',
        'descricao': 'Resumo objetivo da providencia solicitada.',
    },
]

COLUNAS_OBRIGATORIAS_POR_TIPO = [
    {
        'tipo': 'suspensao_liquidacao_reativacao',
        'titulo': 'Suspensao, liquidacao AF e reativacao',
        'campos': [
            {'nome': 'carteira', 'descricao': 'Carteira ou produto: master, pkl_credecesta, asteba_asseba ou avancard.'},
            {'nome': 'cpf', 'descricao': 'CPF do cliente.'},
            {'nome': 'resultado_sentenca', 'descricao': 'Resultado da sentenca.'},
            {'nome': 'data_sentenca', 'descricao': 'Data da sentenca.'},
            {'nome': 'id_oficio', 'descricao': 'ID do oficio relacionado.'},
            {'nome': 'comarca', 'descricao': 'Dados da comarca.'},
        ],
    },
    {
        'tipo': 'cartoes',
        'titulo': 'Cartoes',
        'campos': [
            {'nome': 'carteira', 'descricao': 'Carteira ou produto: master, pkl_credecesta ou investprev_kovr.'},
        ],
    },
]

COLUNAS_OPCIONAIS = [
    {
        'nome': 'numero_processo',
        'descricao': 'Numero do processo, quando quiser levar essa referencia para o email.',
    },
    {
        'nome': 'anexo_pdf',
        'descricao': 'Nome do PDF da linha. Ex.: pedido-123.pdf. Se houver pedido_id, aceita fallback pedido_id.pdf.',
    },
    {
        'nome': 'pedido_id',
        'descricao': 'Codigo interno do pedido. Tambem pode ser usado para localizar o PDF.',
    },
    {
        'nome': 'assunto',
        'descricao': 'Assunto personalizado. Se nao vier, o sistema monta automaticamente.',
    },
    {
        'nome': 'observacoes',
        'descricao': 'Observacoes livres para complementar a solicitacao.',
    },
    {
        'nome': 'email_copia',
        'descricao': 'Emails em copia, separados por virgula ou ponto e virgula.',
    },
    {
        'nome': 'email_destinatario',
        'descricao': 'Campo legado de override manual. O recomendado e deixar o sistema rotear automaticamente.',
    },
]

ROTAS_EMAIL_EXIBICAO = [
    {
        'tipo_label': 'Suspensao, Liquidacao AF e Reativacao',
        'carteira_label': 'Master',
        'destinatario': 'suspensaoeliquidacaobko@bancomaster.com.br',
    },
    {
        'tipo_label': 'Suspensao, Liquidacao AF e Reativacao',
        'carteira_label': 'PKL - Cartao Beneficio (Credecesta)',
        'destinatario': 'desaverbacao@grupoterrafirme.com.br',
    },
    {
        'tipo_label': 'Suspensao e Liquidacao',
        'carteira_label': 'Asteba e Asseba',
        'destinatario': 'desaverbacao@grupoterrafirme.com.br',
    },
    {
        'tipo_label': 'Suspensao e Liquidacao',
        'carteira_label': 'Avancard',
        'destinatario': 'consignado.avancard@bancomaster.com.br',
    },
    {
        'tipo_label': 'Cobranca Consignado',
        'carteira_label': 'Geral',
        'destinatario': 'cobrancaconsignado@bancomaster.com.br',
    },
    {
        'tipo_label': 'Cartoes',
        'carteira_label': 'Master, PKL e Investprev (Kovr)',
        'destinatario': 'backofficecartao@bancomaster.com.br',
    },
]

TIPO_DEMANDA_LABELS = {
    'suspensao_liquidacao_reativacao': 'Suspensao, Liquidacao AF e Reativacao',
    'cobranca_consignado': 'Cobranca Consignado',
    'cartoes': 'Cartoes',
}

TIPO_DEMANDA_ALIASES = {
    'suspensao_liquidacao_reativacao': [
        'suspensao_liquidacao_reativacao',
        'suspensao_liquidacao_af_reativacao',
        'suspensao_liquidacao_af_emprestimo_e_reativacao',
        'suspensao_e_liquidacao',
        'cumprimento_judicial',
        'obrigacao_de_fazer',
        'reativacao',
    ],
    'cobranca_consignado': [
        'cobranca_consignado',
        'cobranca',
        'boleto',
        'emissao_boleto',
        'emissao_de_boletos',
    ],
    'cartoes': [
        'cartoes',
        'cartao',
        'cancelamento_cartao',
        'servicos_cartao',
    ],
}

CARTEIRA_LABELS = {
    'master': 'Master',
    'pkl_credecesta': 'PKL - Cartao Beneficio (Credecesta)',
    'asteba_asseba': 'Asteba e Asseba',
    'avancard': 'Avancard',
    'investprev_kovr': 'Investprev (Kovr)',
}

CARTEIRA_ALIASES = {
    'master': [
        'master',
        'banco_master',
        'bancomaster',
    ],
    'pkl_credecesta': [
        'pkl',
        'credecesta',
        'pkl_credecesta',
        'pkl_cartao_beneficio_credecesta',
        'pkl_cartao_beneficio',
    ],
    'asteba_asseba': [
        'asteba',
        'asseba',
        'asteba_asseba',
        'asteba_e_asseba',
    ],
    'avancard': [
        'avancard',
    ],
    'investprev_kovr': [
        'investprev',
        'kovr',
        'investprev_kovr',
        'atual_kovr',
    ],
}

ROTEAMENTO_MAP = {
    ('suspensao_liquidacao_reativacao', 'master'): 'suspensaoeliquidacaobko@bancomaster.com.br',
    ('suspensao_liquidacao_reativacao', 'pkl_credecesta'): 'desaverbacao@grupoterrafirme.com.br',
    ('suspensao_liquidacao_reativacao', 'asteba_asseba'): 'desaverbacao@grupoterrafirme.com.br',
    ('suspensao_liquidacao_reativacao', 'avancard'): 'consignado.avancard@bancomaster.com.br',
    ('cobranca_consignado', None): 'cobrancaconsignado@bancomaster.com.br',
    ('cartoes', 'master'): 'backofficecartao@bancomaster.com.br',
    ('cartoes', 'pkl_credecesta'): 'backofficecartao@bancomaster.com.br',
    ('cartoes', 'investprev_kovr'): 'backofficecartao@bancomaster.com.br',
}


def _normalizar_texto(valor):
    valor = unicodedata.normalize('NFKD', str(valor or ''))
    valor = valor.encode('ascii', 'ignore').decode('ascii')
    valor = valor.strip().lower()
    valor = re.sub(r'[^a-z0-9]+', '_', valor)
    return valor.strip('_')


def _normalizar_nome_arquivo(nome):
    caminho = Path(str(nome or '').strip())
    base = _normalizar_texto(caminho.stem)
    extensao = caminho.suffix.lower()
    return f"{base}{extensao}" if base or extensao else ''


class MultipleFileInput(forms.ClearableFileInput):
    allow_multiple_selected = True


class MultipleFileField(forms.FileField):
    widget = MultipleFileInput

    def clean(self, data, initial=None):
        if not data:
            return []

        arquivos = data if isinstance(data, (list, tuple)) else [data]
        arquivos_limpos = []
        for arquivo in arquivos:
            arquivos_limpos.append(super().clean(arquivo, initial))
        return arquivos_limpos


_COLUMN_ALIASES = {
    'email_destinatario': [
        'email_destinatario',
        'email',
        'e_mail',
        'destinatario_email',
    ],
    'cliente': [
        'cliente',
        'nome_cliente',
        'cliente_nome',
    ],
    'tipo_demanda': [
        'tipo_demanda',
        'tipo',
        'natureza_demanda',
        'categoria_demanda',
    ],
    'carteira': [
        'carteira',
        'produto',
        'empresa',
        'relacionado_a',
    ],
    'pedido': [
        'pedido',
        'solicitacao',
        'descricao_pedido',
        'demanda',
        'obrigacao',
    ],
    'cpf': [
        'cpf',
        'cpf_cliente',
    ],
    'resultado_sentenca': [
        'resultado_sentenca',
        'resultado_da_sentenca',
        'resultado',
    ],
    'data_sentenca': [
        'data_sentenca',
        'data_da_sentenca',
    ],
    'id_oficio': [
        'id_oficio',
        'oficio_id',
        'id_do_oficio',
    ],
    'comarca': [
        'comarca',
        'dados_da_comarca',
    ],
    'numero_processo': [
        'numero_processo',
        'processo',
        'processo_numero',
        'numero_do_processo',
    ],
    'anexo_pdf': [
        'anexo_pdf',
        'anexo',
        'pdf',
        'arquivo_pdf',
        'nome_pdf',
    ],
    'pedido_id': [
        'pedido_id',
        'id_pedido',
        'codigo_pedido',
        'protocolo',
    ],
    'assunto': [
        'assunto',
        'subject',
    ],
    'observacoes': [
        'observacoes',
        'observacao',
        'obs',
    ],
    'email_copia': [
        'email_copia',
        'email_cc',
        'cc',
        'copia',
    ],
}

_ALIAS_MAP = {}
for _campo, _aliases in _COLUMN_ALIASES.items():
    for _alias in _aliases:
        _ALIAS_MAP[_normalizar_texto(_alias)] = _campo


def _normalizar_celula(valor):
    if valor is None:
        return ''
    return str(valor).strip()


def _resolver_alias(valor, aliases, nome_campo):
    chave = _normalizar_texto(valor)
    for canonico, opcoes in aliases.items():
        if chave in {_normalizar_texto(opcao) for opcao in opcoes}:
            return canonico
    raise ValidationError(f"Valor invalido para {nome_campo}: '{valor}'.")


def _carregar_planilha_csv(conteudo):
    ultimo_erro = None
    texto = None
    for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            texto = conteudo.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            ultimo_erro = exc

    if texto is None:
        raise ValidationError(f"Nao foi possivel ler o CSV enviado: {ultimo_erro}")

    leitor = csv.DictReader(io.StringIO(texto))
    if not leitor.fieldnames:
        raise ValidationError("A planilha CSV precisa ter um cabecalho na primeira linha.")

    return list(leitor.fieldnames), list(leitor)


def _carregar_planilha_xlsx(conteudo):
    workbook = load_workbook(filename=io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        linhas = worksheet.iter_rows(values_only=True)
        try:
            cabecalhos = next(linhas)
        except StopIteration as exc:
            raise ValidationError("A planilha XLSX enviada esta vazia.") from exc

        cabecalhos = [_normalizar_celula(valor) for valor in cabecalhos]
        registros = [dict(zip(cabecalhos, linha)) for linha in linhas]
        return cabecalhos, registros
    finally:
        workbook.close()


def _identificar_colunas(cabecalhos):
    mapa_colunas = {}
    for cabecalho in cabecalhos:
        campo = _ALIAS_MAP.get(_normalizar_texto(cabecalho))
        if campo and campo not in mapa_colunas:
            mapa_colunas[campo] = cabecalho
    return mapa_colunas


def _validar_emails_em_copia(valor, numero_linha):
    emails = [item.strip() for item in re.split(r'[;,]', valor or '') if item.strip()]
    for email in emails:
        try:
            validate_email(email)
        except ValidationError as exc:
            raise ValidationError(f"Linha {numero_linha}: email de copia invalido ({email}).") from exc
    return emails


def _resolver_destinatario(pedido):
    if pedido.get('email_destinatario'):
        try:
            validate_email(pedido['email_destinatario'])
        except ValidationError as exc:
            raise ValidationError(
                f"Linha {pedido['linha_origem']}: email_destinatario invalido."
            ) from exc
        return pedido['email_destinatario']

    tipo_key = pedido['tipo_demanda_key']
    carteira_key = pedido.get('carteira_key')

    destinatario = ROTEAMENTO_MAP.get((tipo_key, carteira_key))
    if destinatario:
        return destinatario

    destinatario = ROTEAMENTO_MAP.get((tipo_key, None))
    if destinatario:
        return destinatario

    raise ValidationError(
        f"Linha {pedido['linha_origem']}: nao foi possivel rotear automaticamente o email para "
        f"tipo_demanda='{pedido.get('tipo_demanda')}' e carteira='{pedido.get('carteira')}'."
    )


def _validar_campos_por_tipo(pedido):
    erros = []
    tipo_key = pedido['tipo_demanda_key']

    if tipo_key == 'suspensao_liquidacao_reativacao':
        obrigatorios = ['carteira', 'cpf', 'resultado_sentenca', 'data_sentenca', 'id_oficio', 'comarca']
    elif tipo_key == 'cartoes':
        obrigatorios = ['carteira']
    else:
        obrigatorios = []

    faltando = [campo for campo in obrigatorios if not pedido.get(campo)]
    if faltando:
        erros.append(
            f"Linha {pedido['linha_origem']}: faltam campos obrigatorios para {TIPO_DEMANDA_LABELS[tipo_key]}: "
            + ", ".join(faltando)
            + "."
        )

    if tipo_key == 'cartoes' and pedido.get('carteira_key') not in {'master', 'pkl_credecesta', 'investprev_kovr'}:
        erros.append(
            f"Linha {pedido['linha_origem']}: a carteira '{pedido.get('carteira')}' nao e valida para demandas de cartoes."
        )

    if tipo_key == 'suspensao_liquidacao_reativacao' and pedido.get('carteira_key') not in {
        'master',
        'pkl_credecesta',
        'asteba_asseba',
        'avancard',
    }:
        erros.append(
            f"Linha {pedido['linha_origem']}: a carteira '{pedido.get('carteira')}' nao e valida para suspensao/liquidacao/reativacao."
        )

    return erros


def _normalizar_regras_negocio(pedido):
    pedido['tipo_demanda_key'] = _resolver_alias(pedido.get('tipo_demanda'), TIPO_DEMANDA_ALIASES, 'tipo_demanda')
    pedido['tipo_demanda_label'] = TIPO_DEMANDA_LABELS[pedido['tipo_demanda_key']]

    if pedido.get('carteira'):
        pedido['carteira_key'] = _resolver_alias(pedido.get('carteira'), CARTEIRA_ALIASES, 'carteira')
        pedido['carteira_label'] = CARTEIRA_LABELS[pedido['carteira_key']]
    else:
        pedido['carteira_key'] = None
        pedido['carteira_label'] = ''

    pedido['email_destinatario'] = _resolver_destinatario(pedido)
    return pedido


def _mapear_anexos_pdf(anexos_pdf):
    anexos_por_nome = {}
    duplicados = []

    for anexo in anexos_pdf or []:
        if Path(anexo.name or '').suffix.lower() != '.pdf':
            raise ValidationError(f"O arquivo '{anexo.name}' nao e um PDF valido.")

        chave = _normalizar_nome_arquivo(anexo.name)
        if chave in anexos_por_nome:
            duplicados.append(Path(anexo.name).name)
            continue

        anexos_por_nome[chave] = {
            'nome': Path(anexo.name).name,
            'conteudo': anexo.read(),
        }
        anexo.seek(0)

    if duplicados:
        raise ValidationError(
            "Existem PDFs com nome repetido no upload: " + ", ".join(sorted(duplicados)) + "."
        )

    return anexos_por_nome


def _vincular_anexos_aos_pedidos(pedidos, anexos_pdf):
    anexos_por_nome = _mapear_anexos_pdf(anexos_pdf)
    erros = []
    anexos_utilizados = set()

    for pedido in pedidos:
        nome_anexo = pedido.get('anexo_pdf', '').strip()
        if not nome_anexo and pedido.get('pedido_id'):
            nome_fallback = f"{pedido['pedido_id']}.pdf"
            chave_fallback = _normalizar_nome_arquivo(nome_fallback)
            if chave_fallback in anexos_por_nome:
                nome_anexo = nome_fallback

        if not nome_anexo:
            pedido['anexo_pdf_nome'] = ''
            pedido['anexo_pdf_conteudo'] = b''
            continue

        chave = _normalizar_nome_arquivo(nome_anexo)
        anexo = anexos_por_nome.get(chave)
        if not anexo:
            erros.append(f"Linha {pedido['linha_origem']}: o PDF '{nome_anexo}' nao foi enviado.")
            continue

        pedido['anexo_pdf_nome'] = anexo['nome']
        pedido['anexo_pdf_conteudo'] = anexo['conteudo']
        anexos_utilizados.add(chave)

    anexos_nao_utilizados = sorted(
        anexo['nome']
        for chave, anexo in anexos_por_nome.items()
        if chave not in anexos_utilizados
    )
    if anexos_nao_utilizados:
        erros.append(
            "PDF(s) enviados sem vinculo na planilha: "
            + ", ".join(anexos_nao_utilizados)
            + ". Use a coluna anexo_pdf ou nomeie o arquivo como pedido_id.pdf."
        )

    if erros:
        raise ValidationError(erros)

    return pedidos


def parsear_planilha_obrigacao_fazer(upload):
    extensao = Path(upload.name or '').suffix.lower()
    conteudo = upload.read()
    upload.seek(0)

    if extensao == '.csv':
        cabecalhos, registros = _carregar_planilha_csv(conteudo)
    elif extensao == '.xlsx':
        cabecalhos, registros = _carregar_planilha_xlsx(conteudo)
    else:
        raise ValidationError("Envie uma planilha no formato .xlsx ou .csv.")

    mapa_colunas = _identificar_colunas(cabecalhos)
    colunas_faltantes = [item['nome'] for item in COLUNAS_OBRIGATORIAS if item['nome'] not in mapa_colunas]
    if colunas_faltantes:
        raise ValidationError("Colunas obrigatorias ausentes: " + ", ".join(colunas_faltantes) + ".")

    pedidos = []
    erros = []

    for indice, registro in enumerate(registros, start=2):
        pedido = {}
        erros_linha = []
        for campo in _COLUMN_ALIASES:
            cabecalho = mapa_colunas.get(campo)
            pedido[campo] = _normalizar_celula(registro.get(cabecalho, '')) if cabecalho else ''

        if not any(pedido.values()):
            continue

        faltando_basico = [item['nome'] for item in COLUNAS_OBRIGATORIAS if not pedido.get(item['nome'])]
        if faltando_basico:
            erros.append(f"Linha {indice}: faltam valores para " + ", ".join(faltando_basico) + ".")
            continue

        try:
            pedido['tipo_demanda_key'] = _resolver_alias(
                pedido.get('tipo_demanda'),
                TIPO_DEMANDA_ALIASES,
                'tipo_demanda',
            )
            pedido['tipo_demanda_label'] = TIPO_DEMANDA_LABELS[pedido['tipo_demanda_key']]
        except ValidationError as exc:
            erros.extend([f"Linha {indice}: {mensagem}" for mensagem in exc.messages])
            continue

        if pedido.get('carteira'):
            try:
                pedido['carteira_key'] = _resolver_alias(pedido.get('carteira'), CARTEIRA_ALIASES, 'carteira')
                pedido['carteira_label'] = CARTEIRA_LABELS[pedido['carteira_key']]
            except ValidationError as exc:
                erros.extend([f"Linha {indice}: {mensagem}" for mensagem in exc.messages])
                continue
        else:
            pedido['carteira_key'] = None
            pedido['carteira_label'] = ''

        pedido['linha_origem'] = indice
        pedido['identificador'] = pedido.get('pedido_id') or pedido.get('cliente') or f"linha {indice}"

        erros_linha.extend(_validar_campos_por_tipo(pedido))
        if erros_linha:
            erros.extend(erros_linha)
            continue

        try:
            pedido['email_destinatario'] = _resolver_destinatario(pedido)
            pedido['cc_list'] = _validar_emails_em_copia(pedido.get('email_copia'), indice)
        except ValidationError as exc:
            erros.extend(exc.messages)
            continue

        pedidos.append(pedido)

    if not pedidos and not erros:
        raise ValidationError("Nenhuma linha valida foi encontrada na planilha enviada.")

    if erros:
        raise ValidationError(erros)

    return pedidos


class DisparoObrigacaoFazerForm(forms.Form):
    planilha = forms.FileField(
        label="Planilha do lote",
        widget=forms.ClearableFileInput(
            attrs={
                'class': 'form-control',
                'accept': '.xlsx,.csv',
            }
        ),
    )
    anexos_pdf = MultipleFileField(
        label="PDFs por pedido",
        required=False,
        widget=MultipleFileInput(
            attrs={
                'class': 'form-control',
                'accept': '.pdf',
                'multiple': True,
            }
        ),
    )
    confirmar_envio = forms.BooleanField(
        label="Confirmo que a planilha foi revisada e que sera enviado 1 email por linha.",
        required=True,
        widget=forms.CheckboxInput(attrs={'class': 'form-check-input'}),
    )

    def clean_anexos_pdf(self):
        anexos = self.cleaned_data.get('anexos_pdf') or []
        for anexo in anexos:
            if Path(anexo.name or '').suffix.lower() != '.pdf':
                raise ValidationError("Todos os anexos enviados precisam estar em PDF.")
        return anexos

    def clean(self):
        cleaned_data = super().clean()
        planilha = cleaned_data.get('planilha')
        if not planilha:
            return cleaned_data

        try:
            pedidos = parsear_planilha_obrigacao_fazer(planilha)
        except ValidationError as exc:
            self.add_error('planilha', exc)
            return cleaned_data

        try:
            cleaned_data['pedidos'] = _vincular_anexos_aos_pedidos(
                pedidos,
                cleaned_data.get('anexos_pdf'),
            )
        except ValidationError as exc:
            self.add_error('anexos_pdf', exc)

        return cleaned_data
